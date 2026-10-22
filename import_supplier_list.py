"""
One-time import of Supplier_List_14_9_26.xlsx into the suppliers table.

Merges into an EXISTING supplier record (by supplier_id, so linked Purchase
Orders/Supplier Invoices stay attached) wherever there's a confident name
match between the existing (abbreviated) name and the new list's fuller
name - rather than creating a duplicate record for the same real company.
Everything else is inserted as a new supplier.

Run once: python3 import_supplier_list.py path/to/Supplier_List.xlsx
"""
import sys
import os
import sqlite3

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import db_location

db_location.migrate_legacy_files(os.path.dirname(os.path.abspath(__file__)))
import openpyxl

DB_PATH = db_location.DB_PATH

# existing supplier_name -> (new fuller name to rename to, or None to keep as-is)
# Only includes matches confident enough to merge rather than risk a wrong
# merge fragmenting or misattributing real purchase history.
MERGE_MAP = {
    "EM Precision": None,
    "FPE Seals": None,
    "Fife Engineering": None,
    "Roemheld": None,
    "Cromwell": "Cromwell Tools",
    "Industrial Interface": "Industrial Interface systems",
    "Pollard": "Pollard Engineering",
    "Quill": "Quill International",
    "Stonebridge": "Stonebridge Testing",
    "Towne": "Towne Lifting",
    "WDS": "WDS Components",
    "Zouch": "Zouch Convertors",
}


def run(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Sheet1"]
    new_rows = {}
    for r in range(3, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        if not name:
            continue
        name = str(name).strip()
        new_rows[name] = (
            ws.cell(row=r, column=3).value,
            ws.cell(row=r, column=4).value,
            ws.cell(row=r, column=5).value,
        )

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    # existing_name -> supplier_id, to find merge targets
    existing = {row["supplier_name"]: row["supplier_id"] for row in conn.execute("SELECT supplier_id, supplier_name FROM suppliers")}
    # new-list name -> existing supplier_id, built from MERGE_MAP
    merge_targets = {}
    for existing_name, new_name in MERGE_MAP.items():
        if existing_name not in existing:
            print(f"WARNING: MERGE_MAP references {existing_name!r} but it's not in the current suppliers table - skipping")
            continue
        lookup_name = new_name or existing_name
        merge_targets[lookup_name] = (existing[existing_name], new_name)

    merged, created, skipped = [], [], []
    for name, (supply_of, reason, auth) in new_rows.items():
        if name in merge_targets:
            supplier_id, rename_to = merge_targets[name]
            if rename_to:
                conn.execute(
                    "UPDATE suppliers SET supplier_name=?, supply_of=?, reason_for_use=?, authorised_by=? WHERE supplier_id=?",
                    (rename_to, supply_of, reason, auth, supplier_id),
                )
                merged.append(f"{name!r} -> merged into existing supplier_id={supplier_id}, renamed to {rename_to!r}")
            else:
                conn.execute(
                    "UPDATE suppliers SET supply_of=?, reason_for_use=?, authorised_by=? WHERE supplier_id=?",
                    (supply_of, reason, auth, supplier_id),
                )
                merged.append(f"{name!r} -> merged into existing supplier_id={supplier_id} (name unchanged)")
        elif name in existing:
            # Exact match not explicitly in MERGE_MAP (shouldn't happen given
            # our analysis, but handle safely rather than creating a dupe)
            conn.execute(
                "UPDATE suppliers SET supply_of=?, reason_for_use=?, authorised_by=? WHERE supplier_id=?",
                (supply_of, reason, auth, existing[name]),
            )
            merged.append(f"{name!r} -> merged into existing supplier_id={existing[name]} (exact name match, auto-detected)")
        else:
            conn.execute(
                "INSERT INTO suppliers (supplier_name, supply_of, reason_for_use, authorised_by) VALUES (?,?,?,?)",
                (name, supply_of, reason, auth),
            )
            created.append(name)

    conn.commit()

    print(f"Merged into existing suppliers: {len(merged)}")
    for m in merged:
        print(" ", m)
    print(f"\nCreated as new suppliers: {len(created)}")
    print(f"\nExisting suppliers not touched (no match in new list): "
          f"{sorted(set(existing) - set(MERGE_MAP.keys()) - set(new_rows.keys()))}")

    total = conn.execute("SELECT COUNT(*) FROM suppliers").fetchone()[0]
    print(f"\nTotal suppliers now: {total}")
    conn.close()


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "data/Supplier_List.xlsx"
    run(path)
